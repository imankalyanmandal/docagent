"""
Regression tests for formatting fidelity and tracked-changes export.

Run from the project root:  python tests/test_fidelity.py

Covers:
  - replace keeps the source's per-word emphasis, on the target's base font
  - editing a caption next to an inline image keeps the image
  - image-only paragraphs are blocks, and move with their section
  - spacing paragraphs travel with (and are deleted with) their block
  - inserting from another document carries images, hyperlinks, list
    numbering and custom styles across correctly
  - tracked-changes export: accept-all == current, reject-all == original
"""
import io
import os
import subprocess
import sys
import tempfile
from copy import deepcopy

sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))
os.environ.setdefault("DOCAGENT_DEBUG", "0")

from docx import Document
from docx.enum.style import WD_STYLE_TYPE
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import qn, nsdecls
from docx.opc.constants import RELATIONSHIP_TYPE as RT
from docx.shared import Pt, Inches
from PIL import Image

import app as appmod
from app import resync, op_replace, op_edit_text, op_move, op_delete, op_insert_after
import tracked

passed = failed = 0


def check(name, cond):
    global passed, failed
    if cond:
        passed += 1
        print(f"  [PASS] {name}")
    else:
        failed += 1
        print(f"  [FAIL] {name}")


def png(color=(200, 30, 30), size=(40, 20)):
    buf = io.BytesIO()
    Image.new("RGB", size, color).save(buf, "PNG")
    buf.seek(0)
    return buf


def entry_of(doc, name="t.docx"):
    buf = io.BytesIO()
    doc.save(buf)
    e = {"filename": name, "document": doc, "undo": [], "original": buf.getvalue()}
    resync(e)
    return e


def find_block(e, pred):
    for bid, b in sorted(e["blocks"].items(), key=lambda kv: kv[1]["order"]):
        if pred(b):
            return bid
    raise AssertionError("block not found")


def reload(doc):
    buf = io.BytesIO()
    doc.save(buf)
    return Document(io.BytesIO(buf.getvalue()))


# ---------------------------------------------------------------------------
print("rich replace")

src = Document()
sp = src.add_paragraph()
sp.add_run("Hello ")
sp.add_run("world").bold = True
sp.add_run(" and ")
r = sp.add_run("more")
r.italic = True
r.font.name = "Comic Sans MS"        # a SOURCE font that must NOT leak in

base = Document()
tp = base.add_paragraph()
tr = tp.add_run("old target text")
tr.font.name = "Arial"
tr.font.size = Pt(14)

be, se = entry_of(base), entry_of(src)
op_replace(be, find_block(be, lambda b: b["type"] == "paragraph"),
           se, find_block(se, lambda b: b["type"] == "paragraph"))
runs = tp.runs
check("replace: text copied", tp.text == "Hello world and more")
check("replace: bold word stays bold",
      [r.text for r in runs if r.bold] == ["world"])
check("replace: italic word stays italic",
      [r.text for r in runs if r.italic] == ["more"])
check("replace: target font/size on every run",
      all(r.font.name == "Arial" and r.font.size == Pt(14) for r in runs))

# table replace keeps per-cell emphasis
src_t = src.add_table(rows=1, cols=2)
src_t.cell(0, 0).paragraphs[0].add_run("plain")
src_t.cell(0, 1).paragraphs[0].add_run("BOLD").bold = True
base_t = base.add_table(rows=2, cols=2)
base_t.cell(0, 0).text = "x"
be, se = entry_of(base), entry_of(src)
op_replace(be, find_block(be, lambda b: b["type"] == "table"),
           se, find_block(se, lambda b: b["type"] == "table"))
check("table replace: shape follows source", len(base_t.rows) == 1)
check("table replace: bold cell stays bold",
      base_t.cell(0, 1).text == "BOLD" and base_t.cell(0, 1).paragraphs[0].runs[0].bold)

# ---------------------------------------------------------------------------
print("images")

d = Document()
d.add_heading("Section A", 1)
d.add_paragraph("A text")
d.add_paragraph().add_run().add_picture(png(), width=Inches(1))      # image-only
d.add_paragraph("")                                                  # spacer
d.add_heading("Section B", 1)
d.add_paragraph("B text")
cap = d.add_paragraph("Figure: ")
cap.add_run().add_picture(png((0, 0, 200)), width=Inches(1))       # caption + image
e = entry_of(d)

img_id = find_block(e, lambda b: b["type"] == "image")
check("image-only paragraph is an image block", e["blocks"][img_id]["images"])
check("spacer attached to the image block", len(e["extras"].get(img_id, [])) == 1)

cap_id = find_block(e, lambda b: b.get("text", "").startswith("Figure"))
op_edit_text(e, cap_id, "Figure 1: the blue box")
resync(e)
cap_el = e["refs"][find_block(e, lambda b: b.get("text", "").startswith("Figure"))]._element
check("edit caption: text updated", e["blocks"][find_block(e, lambda b: b.get("text", "").startswith("Figure"))]["text"] == "Figure 1: the blue box")
check("edit caption: inline image kept", cap_el.find(".//" + qn("w:drawing")) is not None)

# move Section B above Section A - the image must stay inside Section A
op_move(e, find_block(e, lambda b: b.get("text") == "Section B"), "up")
resync(e)
seq = [b["type"] + ":" + (b["text"] or "img") for b in
       sorted(e["blocks"].values(), key=lambda b: b["order"])]
check("move: sections reordered with image inside A",
      seq[:3] == ["header:Section B", "paragraph:B text", "paragraph:Figure 1: the blue box"]
      and seq[3:] == ["header:Section A", "paragraph:A text", "image:img"])
body = list(d.element.body)
img_el = e["refs"][find_block(e, lambda b: b["type"] == "image")]._element
check("move: spacer moved with image",
      body[body.index(img_el) + 1].tag == qn("w:p") and
      not "".join(body[body.index(img_el) + 1].itertext()).strip())

n_before = len(list(d.element.body))
op_delete(e, [find_block(e, lambda b: b["type"] == "image")])
check("delete image block also removes its spacer",
      len(list(d.element.body)) == n_before - 2)

# ---------------------------------------------------------------------------
print("cross-document insert")

other = Document()
other.add_paragraph().add_run().add_picture(png((10, 160, 10)), width=Inches(1))
# hyperlink
hp = other.add_paragraph("See ")
rid = other.part.relate_to("https://example.com/spec", RT.HYPERLINK, is_external=True)
hl = parse_xml(f'<w:hyperlink {nsdecls("w", "r")} r:id="{rid}"><w:r><w:t>the spec</w:t></w:r></w:hyperlink>')
hp._p.append(hl)
# custom style
fancy = other.styles.add_style("Fancy Note", WD_STYLE_TYPE.PARAGRAPH)
fancy.font.color.rgb = None
other.add_paragraph("styled note", style="Fancy Note")
# a list with its OWN numbering definition (id 77) - two items
numbering = other.part.numbering_part.element
numbering.append(parse_xml(
    f'<w:abstractNum {nsdecls("w")} w:abstractNumId="77"><w:nsid w:val="0A0B0C0D"/>'
    f'<w:lvl w:ilvl="0"><w:start w:val="1"/><w:numFmt w:val="upperRoman"/>'
    f'<w:lvlText w:val="%1."/></w:lvl></w:abstractNum>'))
numbering.append(parse_xml(
    f'<w:num {nsdecls("w")} w:numId="77"><w:abstractNumId w:val="77"/></w:num>'))
for txt in ("roman one", "roman two"):
    lp = other.add_paragraph(txt)
    lp._p.get_or_add_pPr().append(parse_xml(
        f'<w:numPr {nsdecls("w")}><w:ilvl w:val="0"/><w:numId w:val="77"/></w:numPr>'))

base = Document()
base.add_paragraph("anchor")
base_ids_before = set(base.part.numbering_part.element.xpath("./w:num/@w:numId"))
be, oe = entry_of(base), entry_of(other)
op_insert_after(be, find_block(be, lambda b: True), oe, list(oe["blocks"]))
base2 = reload(base)          # prove it survives a save/open round trip

blip = base2.element.body.find(".//{http://schemas.openxmlformats.org/drawingml/2006/main}blip")
new_rid = blip.get("{http://schemas.openxmlformats.org/officeDocument/2006/relationships}embed")
src_blob = list(other.part.rels.values())
check("image: relationship resolves in base",
      base2.part.rels[new_rid].reltype == RT.IMAGE)
check("image: same picture bytes", base2.part.rels[new_rid].target_part.blob ==
      next(r.target_part.blob for r in other.part.rels.values() if r.reltype == RT.IMAGE))

hl2 = base2.element.body.find(".//" + qn("w:hyperlink"))
check("hyperlink: points at the same URL in base",
      base2.part.rels[hl2.get(qn("r:id"))].target_ref == "https://example.com/spec")

check("style: custom style copied into base",
      any(s.name == "Fancy Note" for s in base2.styles))

num_ids = [n.get(qn("w:val")) for n in base2.element.body.iter(qn("w:numId"))]
tgt_num = base2.part.numbering_part.element
check("list: both items share one numbering id", len(num_ids) == 2 and len(set(num_ids)) == 1)
num_el = tgt_num.xpath(f'./w:num[@w:numId="{num_ids[0]}"]')
abs_id = num_el[0].xpath("./w:abstractNumId/@w:val")[0] if num_el else None
fmt = tgt_num.xpath(f'./w:abstractNum[@w:abstractNumId="{abs_id}"]/w:lvl/w:numFmt/@w:val')
check("list: numbering definition copied (upper roman)", fmt == ["upperRoman"])
check("list: got a fresh id, not an existing base list", num_ids[0] not in base_ids_before)
order = [c.tag.split("}")[1] for c in tgt_num if c.tag in (qn("w:abstractNum"), qn("w:num"))]
check("list: abstractNums all precede nums (schema order)",
      order == sorted(order, key=lambda t: t != "abstractNum"))

# ---------------------------------------------------------------------------
print("tracked changes")

W = lambda t: qn("w:" + t)


def content(el_list):
    """Comparable view: paragraph texts, table cell texts, image presence."""
    out = []
    for el in el_list:
        if el.tag == W("p"):
            out.append(("p", "".join(t.text or "" for t in el.iter(W("t"))),
                        el.find(".//" + W("drawing")) is not None))
        elif el.tag == W("tbl"):
            out.append(("tbl", tuple(tuple("".join(t.text or "" for t in tc.iter(W("t")))
                                           for tc in tr.findall(W("tc")))
                                     for tr in el.findall(W("tr")))))
    return out


def _paragraph_mark(p, kind):
    return p.find(f"{W('pPr')}/{W('rPr')}/{W(kind)}") is not None


def resolve(doc, accept):
    """Simulate Word's Accept All / Reject All on a copy of the body."""
    body = deepcopy(doc.element.body)
    gone, keep = ("del", "ins") if accept else ("ins", "del")
    for tr in list(body.iter(W("tr"))):
        if tr.find(f"{W('trPr')}/{W(gone)}") is not None:
            tr.getparent().remove(tr)
    for tbl in list(body.iter(W("tbl"))):        # Word drops a table with no rows left
        if tbl.find(W("tr")) is None and tbl.getparent() is not None:
            tbl.getparent().remove(tbl)
    for p in list(body.iter(W("p"))):
        if _paragraph_mark(p, gone) and p.getparent() is not None:
            p.getparent().remove(p)
    for el in list(body.iter(W(gone))):
        if el.getparent() is not None and el.getparent().tag != W("rPr") \
                and el.getparent().tag != W("trPr"):
            el.getparent().remove(el)
    for el in list(body.iter(W(keep))):
        parent = el.getparent()
        if parent.tag in (W("rPr"), W("trPr")):
            parent.remove(el)
            continue
        for c in list(el):
            el.addprevious(c)
        parent.remove(el)
    for t in body.iter(W("delText")):
        t.tag = W("t")
    return content([c for c in body if c.tag != W("sectPr")])


orig = Document()
orig.add_heading("Intro", 1)
orig.add_paragraph("The quick brown fox jumps over the dog.")
orig.add_paragraph("Delete me please.")
t = orig.add_table(rows=2, cols=2)
for (rr, cc), v in {(0, 0): "Name", (0, 1): "Qty", (1, 0): "Apple", (1, 1): "3"}.items():
    t.cell(rr, cc).text = v
orig.add_paragraph().add_run().add_picture(png(), width=Inches(1))
orig.add_heading("End", 2)
p_last = orig.add_paragraph("Final ")
p_last.add_run("words").bold = True

client = appmod.app.test_client()
buf = io.BytesIO()
orig.save(buf)
orig_bytes = buf.getvalue()
buf.seek(0)
doc_id = client.post("/upload", data={"file": (buf, "orig.docx")},
                     content_type="multipart/form-data").get_json()["doc_id"]
E = appmod.DOC_STORE[doc_id]
bid = lambda pred: find_block(E, pred)
post = lambda route, **kw: client.post(route, json={"doc_id": doc_id, **kw})

post("/edit_text", block_id=bid(lambda b: b.get("text", "").startswith("The quick")),
     text="The quick red fox leaps over the dog.")
post("/delete", block_ids=[bid(lambda b: b.get("text") == "Delete me please.")], cascade=False)
post("/table", block_id=bid(lambda b: b["type"] == "table"), operation="set_content",
     rows=[["Name", "Qty"], ["Apple", "5"], ["Pear", "2"]])
post("/delete", block_ids=[bid(lambda b: b["type"] == "image")], cascade=False)
post("/find_replace", find="Final", replace="Closing")
end_h = bid(lambda b: b.get("text") == "End")
appmod.set_heading_level(E["refs"][end_h], 1)        # style change
resync(E)

r = client.get(f"/export/{doc_id}?track=1&author=Reviewer")
check("tracked export: 200 + docx", r.status_code == 200 and r.data[:2] == b"PK")
out = Document(io.BytesIO(r.data))
check("tracked export: has insertions and deletions",
      out.element.body.find(".//" + W("ins")) is not None and
      out.element.body.find(".//" + W("del")) is not None)
check("tracked export: author recorded",
      out.element.body.find(".//" + W("ins")).get(W("author")) == "Reviewer")

current = content([c for c in E["document"].element.body if c.tag != W("sectPr")])
original = content([c for c in Document(io.BytesIO(orig_bytes)).element.body
                    if c.tag != W("sectPr")])
check("tracked export: Accept All == current document", resolve(out, True) == current)
check("tracked export: Reject All == original document", resolve(out, False) == original)

dels = "".join(t.text for t in out.element.body.iter(W("delText")))
check("tracked export: word-level diff (only changed words struck)",
      "brown" in dels and "jumps" in dels and "quick" not in dels)
check("tracked export: style change recorded",
      out.element.body.find(".//" + W("pPrChange")) is not None)
ids = [el.get(W("id")) for tag in ("ins", "del", "pPrChange") for el in out.element.body.iter(W(tag))]
check("tracked export: revision ids unique", len(ids) == len(set(ids)))
check("tracked export: live document untouched",
      E["document"].element.body.find(".//" + W("ins")) is None)

plain = client.get(f"/export/{doc_id}")
check("plain export still has no revisions",
      Document(io.BytesIO(plain.data)).element.body.find(".//" + W("ins")) is None)

# LibreOffice (if installed) must be able to open both exports
soffice = next((p for p in ("/usr/bin/soffice", "/usr/bin/libreoffice") if os.path.exists(p)), None)
if soffice:
    with tempfile.TemporaryDirectory() as tmp:
        for name, data in (("tracked.docx", r.data), ("cross.docx", None)):
            path = os.path.join(tmp, name)
            if data is None:
                base2.save(path)
            else:
                open(path, "wb").write(data)
            subprocess.run([soffice, "--headless", "--convert-to", "pdf", "--outdir", tmp, path],
                           capture_output=True, timeout=120)
            check(f"LibreOffice opens {name}",
                  os.path.exists(os.path.join(tmp, name.replace(".docx", ".pdf"))))
else:
    print("  [SKIP] LibreOffice not installed - skipping open-in-office check")

print(f"\n=== {passed} passed, {failed} failed ===")
sys.exit(1 if failed else 0)
