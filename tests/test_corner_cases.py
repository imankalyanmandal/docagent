"""
General corner cases: moving, deleting, duplicating, undo, find & replace,
text editing, cross-document insert and tracked-changes export at their
edges.

Run from the project root:  python tests/test_corner_cases.py
"""
import io
import os
import sys

sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))
os.environ.setdefault("DOCAGENT_DEBUG", "0")

from docx import Document
from docx.oxml.ns import qn
from docx.opc.constants import RELATIONSHIP_TYPE as RT
from docx.shared import Inches
from PIL import Image

import app as appmod

passed = failed = 0
W = lambda t: qn("w:" + t)
client = appmod.app.test_client()


def check(name, cond, detail=""):
    global passed, failed
    if cond:
        passed += 1
        print(f"  [PASS] {name}")
    else:
        failed += 1
        print(f"  [FAIL] {name}" + (f"  -> {detail}" if detail else ""))


def section(t):
    print(f"\n--- {t}")


def png(color=(200, 30, 30)):
    b = io.BytesIO()
    Image.new("RGB", (30, 15), color).save(b, "PNG")
    b.seek(0)
    return b


def upload(doc, name="t.docx"):
    buf = io.BytesIO()
    doc.save(buf)
    buf.seek(0)
    r = client.post("/upload", data={"file": (buf, name)}, content_type="multipart/form-data")
    assert r.status_code == 200, r.get_json()
    return r.get_json()["doc_id"]


def call(route, **payload):
    r = client.post(route, json=payload)
    return r.status_code, (r.get_json() or {})


def blocks(doc_id):
    return sorted(appmod.DOC_STORE[doc_id]["blocks"].values(), key=lambda b: b["order"])


def texts(doc_id):
    return [b.get("text") if b["type"] != "table" else "<table>" for b in blocks(doc_id)]


def bid(doc_id, text):
    return next(b["id"] for b in blocks(doc_id) if b.get("text") == text)


def outline_doc():
    d = Document()
    d.add_heading("A", 1)
    d.add_paragraph("a1")
    d.add_heading("A.1", 2)
    d.add_paragraph("a1x")
    d.add_heading("A.2", 2)
    d.add_paragraph("a2x")
    d.add_heading("B", 1)
    d.add_paragraph("b1")
    return d


def body_xml(doc_id):
    buf = io.BytesIO()
    appmod.DOC_STORE[doc_id]["document"].save(buf)
    return buf.getvalue()


# ===========================================================================
section("ids and hierarchy")

d = Document()
d.add_paragraph("preamble before any heading")
d.add_heading("Top", 1)
d.add_heading("Skipped to level 3", 3)
d.add_paragraph("deep")
d.add_heading("Title style", 0)          # "Title" counts as level 1
d.add_paragraph("   ")                   # whitespace-only paragraph: not a block
d.add_heading("   ", 2)                  # whitespace-only heading: not a block
doc_id = upload(d)
ids = [(b["id"], b.get("text")) for b in blocks(doc_id)]
check("preamble before the first heading gets a top-level id", ids[0] == ("1", "preamble before any heading"), ids)
check("skipped heading level still nests under its parent", ("2.1", "Skipped to level 3") in ids, ids)
check("paragraph under a level-3 heading nests below it", ("2.1.1", "deep") in ids, ids)
check("Title style is a level-1 heading", any(t == "Title style" and i.count(".") == 0 for i, t in ids), ids)
check("whitespace-only paragraph/heading are not blocks", all((t or "x").strip() for _, t in ids), ids)

empty = upload(Document())
check("empty document loads with no blocks", blocks(empty) == [])
status, body = call("/undo", doc_id=empty)
check("undo on a fresh document refused", status == 400, body)

# ===========================================================================
section("move")

doc_id = upload(outline_doc())
status, _ = call("/move", doc_id=doc_id, block_id=bid(doc_id, "A"), direction="up")
check("moving the first block up is a harmless no-op", status in (200, 400) and texts(doc_id)[0] == "A")
status, _ = call("/move", doc_id=doc_id, block_id=bid(doc_id, "b1"), direction="down")
check("moving the last block down is a harmless no-op", texts(doc_id)[-1] == "b1")
status, body = call("/move", doc_id=doc_id, block_id=bid(doc_id, "A"), direction="sideways")
check("invalid direction refused", status == 400, body)
status, body = call("/move", doc_id=doc_id, block_id="9.9.9", direction="up")
check("unknown block refused", status == 400, body)

call("/move", doc_id=doc_id, block_id=bid(doc_id, "A.2"), direction="up")
check("move a subsection up swaps it with its sibling (children follow)",
      texts(doc_id) == ["A", "a1", "A.2", "a2x", "A.1", "a1x", "B", "b1"], texts(doc_id))
call("/move", doc_id=doc_id, block_id=bid(doc_id, "B"), direction="up")
check("move a top-level section up past a section with subsections",
      texts(doc_id) == ["B", "b1", "A", "a1", "A.2", "a2x", "A.1", "a1x"], texts(doc_id))

status, body = call("/move_to", doc_id=doc_id, block_ids=[bid(doc_id, "A")],
                    anchor_block_id=bid(doc_id, "a1x"), position="after")
check("move_to an anchor inside the moved section refused", status == 400, body)
call("/move_to", doc_id=doc_id, block_ids=[bid(doc_id, "b1"), bid(doc_id, "a1")],
     anchor_block_id=bid(doc_id, "A.1"), position="after")
check("move_to with several blocks keeps their document order, lands after the anchor's section",
      texts(doc_id) == ["B", "A", "A.2", "a2x", "A.1", "a1x", "b1", "a1"], texts(doc_id))
call("/move_to", doc_id=doc_id, block_ids=[bid(doc_id, "a1")],
     anchor_block_id=bid(doc_id, "B"), position="before")
check("move_to before the very first block", texts(doc_id)[0] == "a1", texts(doc_id))

# ===========================================================================
section("delete")

doc_id = upload(outline_doc())
call("/delete", doc_id=doc_id, block_ids=[bid(doc_id, "A")], cascade=False)
check("delete heading without cascade keeps its content", "a1" in texts(doc_id) and "A" not in texts(doc_id))
ids = {b.get("text"): b["id"] for b in blocks(doc_id)}
check("orphaned subsections re-number cleanly", ids["A.1"] == "1" or ids["A.1"].count(".") <= 1, ids)
call("/delete", doc_id=doc_id, block_ids=[bid(doc_id, "B")], cascade=True)
check("cascade delete removes the section", "b1" not in texts(doc_id))
call("/delete", doc_id=doc_id, block_ids=[b["id"] for b in blocks(doc_id)], cascade=True)
check("deleting every block leaves an empty (valid) document",
      blocks(doc_id) == [] and Document(io.BytesIO(body_xml(doc_id))) is not None)
status, body = call("/delete", doc_id=doc_id, block_ids=["1"], cascade=True)
check("delete an unknown block refused", status == 400, body)
status, body = call("/delete", doc_id=doc_id, block_ids=[], cascade=True)
check("delete nothing does not crash", status in (200, 400), body)

# duplicate ids in the request
doc_id = upload(outline_doc())
status, _ = call("/delete", doc_id=doc_id, block_ids=[bid(doc_id, "a1"), bid(doc_id, "a1")], cascade=True)
check("same id twice in a delete request is fine", status == 200 and "a1" not in texts(doc_id))
# parent and child both selected
status, _ = call("/delete", doc_id=doc_id, block_ids=[bid(doc_id, "A"), bid(doc_id, "A.1")], cascade=True)
check("parent + child both selected deletes once without error", status == 200 and "a1x" not in texts(doc_id))

# ===========================================================================
section("duplicate")

d = outline_doc()
d.paragraphs[1].add_run().add_picture(png(), width=Inches(0.5))
b = d.add_paragraph()
b._p.append(d.element.makeelement(W("bookmarkStart"), {W("id"): "7", W("name"): "keep_me"}))
b.add_run("bookmarked")
b._p.append(d.element.makeelement(W("bookmarkEnd"), {W("id"): "7"}))
doc_id = upload(d)
call("/duplicate", doc_id=doc_id, block_ids=[bid(doc_id, "A")])
check("duplicate a section clones heading + all nested content",
      texts(doc_id)[:12] == ["A", "a1", "A.1", "a1x", "A.2", "a2x"] * 2, texts(doc_id))
check("duplicated section's image points at a real image",
      all(appmod.DOC_STORE[doc_id]["document"].part.rels[b.get(qn("r:embed"))].reltype == RT.IMAGE
          for b in appmod.DOC_STORE[doc_id]["document"].element.body.iter(
              "{http://schemas.openxmlformats.org/drawingml/2006/main}blip")))
call("/duplicate", doc_id=doc_id, block_ids=[bid(doc_id, "bookmarked")])
names = [e.get(W("name")) for e in appmod.DOC_STORE[doc_id]["document"].element.body.iter(W("bookmarkStart"))]
check("duplicating a bookmarked paragraph doesn't duplicate the bookmark name",
      names.count("keep_me") == 1, names)

# ===========================================================================
section("undo")

doc_id = upload(outline_doc())
original = texts(doc_id)
status, _ = call("/move", doc_id=doc_id, block_id="9", direction="up")      # fails
check("a failed operation leaves nothing to undo", call("/undo", doc_id=doc_id)[0] == 400)
for i in range(25):
    call("/edit_text", doc_id=doc_id, block_id=bid(doc_id, "b1") if i == 0 else bid(doc_id, f"b1-{i - 1}"),
         text=f"b1-{i}")
undos = 0
while call("/undo", doc_id=doc_id)[0] == 200:
    undos += 1
check("undo is capped at 20 steps", undos == 20, undos)
check("after 20 undos the doc is 5 edits in, not corrupted", "b1-4" in texts(doc_id), texts(doc_id))

doc_id = upload(outline_doc())
call("/delete", doc_id=doc_id, block_ids=[bid(doc_id, "A")], cascade=True)
call("/duplicate", doc_id=doc_id, block_ids=[bid(doc_id, "B")])
call("/undo", doc_id=doc_id)
call("/undo", doc_id=doc_id)
check("mixed operations undo back to the exact original", texts(doc_id) == original, texts(doc_id))

# ===========================================================================
section("edit text")

doc_id = upload(outline_doc())
status, _ = call("/edit_text", doc_id=doc_id, block_id=bid(doc_id, "a1"), text="  spaced  ")
check("edit keeps leading/trailing spaces", status == 200 and "  spaced  " in texts(doc_id))
import zipfile
xml = zipfile.ZipFile(io.BytesIO(body_xml(doc_id))).read("word/document.xml")
check("leading/trailing spaces are preserved in the saved file (xml:space)",
      b'xml:space="preserve">  spaced  <' in xml)
status, _ = call("/edit_text", doc_id=doc_id, block_id=bid(doc_id, "a2x"), text="tab\there\nnew line")
check("tab and line break round-trip", "tab\there\nnew line" in texts(doc_id), texts(doc_id))
status, _ = call("/edit_text", doc_id=doc_id, block_id=bid(doc_id, "a1x"), text="<b>&amp; \"quotes\"</b>")
check("markup-looking text stored literally", '<b>&amp; "quotes"</b>' in texts(doc_id))
status, _ = call("/edit_text", doc_id=doc_id, block_id=bid(doc_id, "A.1"), text="Renamed heading")
check("editing a heading keeps it a heading",
      any(b["type"] == "header" and b["text"] == "Renamed heading" for b in blocks(doc_id)))
n = len(blocks(doc_id))
status, _ = call("/edit_text", doc_id=doc_id, block_id=bid(doc_id, "b1"), text="")
check("editing text to empty removes the block from the list (spacing paragraph)",
      status == 200 and len(blocks(doc_id)) == n - 1)
d = Document()
d.add_paragraph().add_run().add_picture(png(), width=Inches(0.5))
iid = upload(d)
status, body = call("/edit_text", doc_id=iid, block_id=blocks(iid)[0]["id"], text="x")
check("text-editing an image block refused", status == 400, body)

# hyperlink text is replaced, not duplicated
d = Document()
p = d.add_paragraph("See ")
rid = d.part.relate_to("https://example.com", RT.HYPERLINK, is_external=True)
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls
p._p.append(parse_xml(f'<w:hyperlink {nsdecls("w", "r")} r:id="{rid}"><w:r><w:t>the link</w:t></w:r></w:hyperlink>'))
hid = upload(d)
call("/edit_text", doc_id=hid, block_id=blocks(hid)[0]["id"], text="Replaced entirely")
check("editing a paragraph with a hyperlink doesn't duplicate the link text",
      texts(hid) == ["Replaced entirely"], texts(hid))

# ===========================================================================
section("find & replace")

d = Document()
p = d.add_paragraph()
for chunk in ("Sa", "fe", "ty ", "first. "):
    p.add_run(chunk)
p.add_run("Safety").bold = True
d.add_paragraph("a.b* literal (regex chars) a.b*")
d.add_paragraph("ÉCOLE école École")
d.add_paragraph("aaa")
t = d.add_table(rows=1, cols=2)
t.cell(0, 0).text = "Safety in table"
fid = upload(d)

status, body = call("/find_replace", doc_id=fid, find="Safety", replace="Security")
check("match split across 3 runs + a bold run: both replaced", body.get("replacements") == 3, body)
para = appmod.DOC_STORE[fid]["refs"][blocks(fid)[0]["id"]]
check("replacement inside the bold run stays bold",
      any(r.bold and r.text == "Security" for r in para.runs), [(r.text, r.bold) for r in para.runs])
check("text around a cross-run match intact", blocks(fid)[0]["text"] == "Security first. Security")

status, body = call("/find_replace", doc_id=fid, find="a.b*", replace="X")
check("regex special characters are matched literally", body.get("replacements") == 2
      and "X literal (regex chars) X" in texts(fid), body)
status, body = call("/find_replace", doc_id=fid, find="école", replace="school")
check("case-insensitive matches accented capitals", body.get("replacements") == 3, body)
status, body = call("/find_replace", doc_id=fid, find="aa", replace="b")
check("non-overlapping matches ('aaa' has one 'aa')", body.get("replacements") == 1 and "ba" in texts(fid), body)
status, body = call("/find_replace", doc_id=fid, find="Security", replace="", scope="text")
check("replace with empty string deletes the text", body.get("replacements") == 2
      and blocks(fid)[0]["text"] == " first. ", (body, blocks(fid)[0]))
check("scope=text leaves tables alone",
      appmod.DOC_STORE[fid]["blocks"][next(b["id"] for b in blocks(fid) if b["type"] == "table")]["rows"][0][0]
      == "Security in table")
status, body = call("/find_replace", doc_id=fid, find="", replace="x")
check("empty search refused", status == 400, body)
status, body = call("/find_replace", doc_id=fid, find="zzz-not-there", replace="x")
check("no matches -> 0 replacements", body.get("replacements") == 0, body)

# ===========================================================================
section("cross-document insert")

src = Document()
for _ in range(3):
    src.add_paragraph().add_run().add_picture(png((0, 0, 255)), width=Inches(0.5))   # same image x3
src_id = upload(src, "src.docx")
dst_id = upload(outline_doc(), "dst.docx")
call("/insert", base_doc_id=dst_id, anchor_block_id=bid(dst_id, "b1"),
     source_doc_id=src_id, source_block_ids=[b["id"] for b in blocks(src_id)])
dst_part = appmod.DOC_STORE[dst_id]["document"].part
image_parts = {r.target_part.partname for r in dst_part.rels.values() if r.reltype == RT.IMAGE}
check("inserting the same image 3 times stores it once", len(image_parts) == 1, image_parts)
check("all three image blocks inserted", sum(b["type"] == "image" for b in blocks(dst_id)) == 3)

status, body = call("/insert", base_doc_id=dst_id, anchor_block_id=bid(dst_id, "A"),
                    source_doc_id=dst_id, source_block_ids=[bid(dst_id, "b1")])
check("insert from the same document (self) works", status == 200 and texts(dst_id).count("b1") == 2, body)
status, body = call("/insert", base_doc_id=dst_id, anchor_block_id=bid(dst_id, "A"),
                    source_doc_id=src_id, source_block_ids=[])
check("insert with no source blocks refused", status == 400, body)
status, body = call("/replace", base_doc_id=dst_id, base_block_id=bid(dst_id, "A"),
                    source_doc_id=dst_id, source_block_id=bid(dst_id, "A"))
check("replace a block with itself is a harmless no-op", status == 200 and "A" in texts(dst_id), body)

# heading <-> paragraph replace keeps block types sensible
status, _ = call("/replace", base_doc_id=dst_id, base_block_id=bid(dst_id, "a1"),
                 source_doc_id=dst_id, source_block_id=bid(dst_id, "B"))
check("paragraph replaced by a heading becomes a heading",
      sum(1 for b in blocks(dst_id) if b["type"] == "header" and b["text"] == "B") == 2)

# ===========================================================================
section("tracked-changes export edges")

doc_id = upload(outline_doc())
r = client.get(f"/export/{doc_id}?track=1")
out = Document(io.BytesIO(r.data))
check("no edits -> tracked export has no revisions",
      out.element.body.find(".//" + W("ins")) is None and out.element.body.find(".//" + W("del")) is None)
call("/edit_text", doc_id=doc_id, block_id=bid(doc_id, "b1"), text="changed")
call("/undo", doc_id=doc_id)
out = Document(io.BytesIO(client.get(f"/export/{doc_id}?track=1").data))
check("edit then undo -> no revisions", out.element.body.find(".//" + W("ins")) is None)

call("/delete", doc_id=doc_id, block_ids=[b["id"] for b in blocks(doc_id)], cascade=True)
out = Document(io.BytesIO(client.get(f"/export/{doc_id}?track=1").data))
check("everything deleted -> every original paragraph shows as deleted",
      len(list(out.element.body.iter(W("delText")))) >= 8 and out.element.body.find(".//" + W("t")) is None)

src_id = upload(outline_doc())
call("/upload_html", html="<h1>Fresh</h1><p>all new</p>", append_to=src_id)
out = Document(io.BytesIO(client.get(f"/export/{src_id}?track=1").data))
ins_text = "".join(t.text for el in out.element.body.iter(W("ins")) for t in el.iter(W("t")))
check("appended HTML shows as inserted", "Fresh" in ins_text and "all new" in ins_text, ins_text)

r = client.get(f"/export/{src_id}?track=1&author=" + "A" * 500)
au = Document(io.BytesIO(r.data)).element.body.find(".//" + W("ins")).get(W("author"))
check("very long author name is truncated", len(au) <= 64, len(au))
r = client.get(f"/export/{src_id}?track=1&author=%20%20")
au = Document(io.BytesIO(r.data)).element.body.find(".//" + W("ins")).get(W("author"))
check("blank author falls back to 'docagent'", au == "docagent", au)
check("unknown doc export -> 400", client.get("/export/nope?track=1").status_code == 400)

# ===========================================================================
section("request validation")

status, body = call("/edit_text", doc_id="nope", block_id="1", text="x")
check("unknown doc_id refused", status == 400, body)
r = client.post("/edit_text", data="not json", content_type="application/json")
check("malformed JSON body answered with a clean error (no crash page)",
      r.status_code in (400, 500) and r.is_json, r.status_code)
status, body = call("/table", doc_id=doc_id, block_id="1", operation="explode")
check("unknown table operation refused", status == 400, body)
r = client.post("/upload", data={"file": (io.BytesIO(b"not a docx"), "fake.docx")},
                content_type="multipart/form-data")
check("corrupt .docx upload refused cleanly", r.status_code == 400 and r.is_json)
r = client.post("/upload", data={"file": (io.BytesIO(b"x"), "notes.txt")},
                content_type="multipart/form-data")
check("non-.docx upload refused", r.status_code == 400)
check("unknown media 404s", client.get(f"/media/{src_id}/rId9999").status_code == 404)

print(f"\n=== {passed} passed, {failed} failed ===")
sys.exit(1 if failed else 0)
