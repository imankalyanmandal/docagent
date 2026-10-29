"""
Regression tests for find & replace, request locking, upload limits and
server defaults.

Run from the project root:  python tests/test_backend_fixes.py

Guards these bugs:
  - a paragraph with one match inside a run AND one spanning two runs lost
    the spanning match (and reported the wrong count)
  - a replacement containing a backslash crashed ("bad escape")
  - spanning matches flattened the whole paragraph into one run, dropping
    bold/italic on unrelated words
  - merged table cells were replaced once per spanned row
  - concurrent requests on the same document weren't serialized
  - oversized uploads weren't rejected
"""
import io
import os
import sys
import threading

sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))
os.environ.setdefault("DOCAGENT_DEBUG", "0")

from docx import Document
import app as appmod
from app import _replace_in_paragraph, op_find_replace, STORE_LOCK

passed = 0
failed = 0


def check(name, cond):
    global passed, failed
    if cond:
        passed += 1
        print(f"  [PASS] {name}")
    else:
        failed += 1
        print(f"  [FAIL] {name}")


def para(*runs):
    """Build a paragraph from (text, bold) tuples or plain strings."""
    d = Document()
    p = d.add_paragraph()
    for r in runs:
        text, bold = (r, False) if isinstance(r, str) else r
        p.add_run(text).bold = bold
    return p


print("find & replace")

# 1. mixed: one in-run match + one spanning match in the same paragraph
p = para("cat and ca", "t")
n = _replace_in_paragraph(p, "cat", "dog", False)
check("mixed in-run + spanning: both replaced", p.text == "dog and dog")
check("mixed in-run + spanning: count is 2", n == 2)

# 2. backslashes in replacement are literal
p = para("path here")
try:
    _replace_in_paragraph(p, "path", r"C:\docs\1", False)
    check("backslash replacement is literal", p.text == r"C:\docs\1 here")
except Exception as e:
    check(f"backslash replacement raised {e!r}", False)

# 3. spanning match keeps formatting of unrelated runs
p = para(("Important", True), " note: ca", "t food")
_replace_in_paragraph(p, "cat", "dog", False)
check("spanning match: text correct", p.text == "Important note: dog food")
check("spanning match: bold run untouched",
      p.runs[0].text == "Important" and p.runs[0].bold is True)

# 4. match spanning three runs
p = para("ab", "c", "de")
n = _replace_in_paragraph(p, "bcd", "X", False)
check("three-run span", p.text == "aXe" and n == 1)

# 5. replacement containing the search text doesn't compound
p = para("a b a")
n = _replace_in_paragraph(p, "a", "aa", False)
check("replacement containing find text", p.text == "aa b aa" and n == 2)

# 6. case-insensitive / case-sensitive
p = para("Cat cat CAT")
check("case-insensitive count", _replace_in_paragraph(p, "cat", "x", False) == 3)
p = para("Cat cat CAT")
check("case-sensitive count", _replace_in_paragraph(p, "cat", "x", True) == 1
      and p.text == "Cat x CAT")

# 7. no match leaves paragraph alone
p = para(("keep", True), " me")
check("no match returns 0", _replace_in_paragraph(p, "zzz", "y", False) == 0
      and p.runs[0].bold is True and p.text == "keep me")

# 8. merged table cell replaced once, not once per spanned row
d = Document()
t = d.add_table(rows=3, cols=2)
merged = t.cell(0, 0).merge(t.cell(2, 0))
merged.text = "a"
t.cell(0, 1).text = "a"
entry = {"filename": "t.docx", "document": d, "undo": []}
total = op_find_replace(entry, "a", "aa")
check("merged cell replaced once", t.cell(0, 0).text == "aa")
check("merged cell count correct", total == 2)

print("server")

client = appmod.app.test_client()


def upload(doc):
    buf = io.BytesIO()
    doc.save(buf)
    buf.seek(0)
    return client.post("/upload", data={"file": (buf, "t.docx")},
                       content_type="multipart/form-data")


d = Document()
d.add_paragraph("hello world")
r = upload(d)
check("upload works", r.status_code == 200)
doc_id = r.get_json()["doc_id"]

r = client.post("/find_replace", json={"doc_id": doc_id, "find": "world",
                                       "replace": r"C:\x"})
check("/find_replace with backslash returns 200", r.status_code == 200)

# concurrent edits are serialized: hold the lock, fire a request from another
# thread, and confirm it can't complete until the lock is released
done = threading.Event()


def worker():
    client.post("/edit_text", json={"doc_id": doc_id, "block_id": "1", "text": "x"})
    done.set()


with STORE_LOCK:
    th = threading.Thread(target=worker)
    th.start()
    check("request waits while document is locked", not done.wait(0.5))
th.join(5)
check("request completes after lock is released", done.is_set())

# oversized upload rejected with a JSON 413
old = appmod.app.config["MAX_CONTENT_LENGTH"]
appmod.app.config["MAX_CONTENT_LENGTH"] = 1024
big = io.BytesIO(b"x" * 4096)
r = client.post("/upload", data={"file": (big, "big.docx")},
                content_type="multipart/form-data")
appmod.app.config["MAX_CONTENT_LENGTH"] = old
check("oversized upload -> 413", r.status_code == 413)
check("413 body is JSON", (r.get_json() or {}).get("error", "").startswith("file too large"))

print(f"\n=== {passed} passed, {failed} failed ===")
sys.exit(1 if failed else 0)
