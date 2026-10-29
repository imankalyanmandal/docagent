"""
Regenerate fixtures.json used by tests/ui_test.mjs.

The UI test runs in jsdom (no real browser, no live server), so it needs
canned backend responses that have the EXACT shape the Flask app returns.
Rather than hand-write those, we capture them from the real app via Flask's
test client. Re-run this whenever the backend response shape changes:

    python make_fixtures.py
"""
import json

from docx import Document
from app import app


def build_single_table_doc(path):
    doc = Document()
    doc.add_heading("Report", level=1)
    doc.add_paragraph("Intro paragraph.")
    doc.add_heading("Data", level=2)
    table = doc.add_table(rows=3, cols=3)
    table.style = "Light Grid Accent 1"
    rows = [["Name", "Dept", "Score"], ["Ann", "Eng", "10"], ["Bob", "Ops", "20"]]
    for r, row in enumerate(rows):
        for c, val in enumerate(row):
            table.cell(r, c).text = val
    doc.save(path)


def build_two_table_doc(path):
    """Two tables with a shared header row, for the merge-popup fixtures."""
    doc = Document()
    doc.add_heading("Report", level=1)
    doc.add_paragraph("Intro paragraph.")
    doc.add_heading("Data", level=2)

    t1 = doc.add_table(rows=2, cols=2)
    t1.style = "Light Grid Accent 1"
    for r, row in enumerate([["N", "V"], ["a", "1"]]):
        for c, val in enumerate(row):
            t1.cell(r, c).text = val

    doc.add_paragraph("Between.")

    t2 = doc.add_table(rows=2, cols=2)
    t2.style = "Medium Shading 1 Accent 2"
    for r, row in enumerate([["N", "V"], ["b", "2"]]):
        for c, val in enumerate(row):
            t2.cell(r, c).text = val

    doc.save(path)


def capture():
    client = app.test_client()
    fixtures = {}

    def upload(path):
        with open(path, "rb") as f:
            return client.post(
                "/upload", data={"file": (f, path)}, content_type="multipart/form-data"
            ).get_json()

    build_single_table_doc("uitest.docx")
    build_two_table_doc("uitest_merge.docx")

    up = upload("uitest.docx")
    fixtures["upload"] = up
    fixtures["upload_merge"] = upload("uitest_merge.docx")

    # one fresh document per table op so they're independent
    for name, payload in [
        ("add_row", {"operation": "add_row"}),
        ("add_column", {"operation": "add_column", "header": "Status"}),
        ("delete_column", {"operation": "delete_column", "index": 1}),
        ("delete_row", {"operation": "delete_row", "index": 1}),
        ("sort", {"operation": "sort", "column": 2}),
        ("transpose", {"operation": "transpose"}),
    ]:
        u = upload("uitest.docx")
        tid = next(k for k, v in u["blocks"].items() if v["type"] == "table")
        r = client.post("/table", json={"doc_id": u["doc_id"], "block_id": tid, **payload})
        fixtures[name] = {
            "status": r.status_code,
            "body": r.get_json(),
            "request": {"doc_id": u["doc_id"], "block_id": tid, **payload},
        }

    # set_content - the table editor popup's Save action
    u = upload("uitest.docx")
    tid = next(k for k, v in u["blocks"].items() if v["type"] == "table")
    new_rows = [["Name", "Dept", "Score"], ["Zed", "Ops", "99"], ["Ann", "Eng", "10"]]
    r = client.post(
        "/table",
        json={"doc_id": u["doc_id"], "block_id": tid, "operation": "set_content", "rows": new_rows},
    )
    fixtures["set_content"] = {
        "status": r.status_code, "body": r.get_json(),
        "request": {"doc_id": u["doc_id"], "block_id": tid},
    }

    # edit_text - inline paragraph/header editing
    u = upload("uitest.docx")
    r = client.post("/edit_text", json={"doc_id": u["doc_id"], "block_id": "1.1", "text": "Edited intro text."})
    fixtures["edit_text"] = {
        "status": r.status_code, "body": r.get_json(),
        "request": {"doc_id": u["doc_id"], "block_id": "1.1"},
    }

    # upload_html - Froala paste converted to a new document
    froala = ("<h1>Pasted Report</h1><p>Some <strong>bold</strong> text.</p>"
              "<p class='caption'>Table 1: metrics</p>"
              "<table><tr><th>N</th><th>V</th></tr><tr><td>a</td><td>1</td></tr></table>"
              "<ul><li>one</li><li>two</li></ul>")
    r = client.post("/upload_html", json={"html": froala, "name": "froala.html"})
    fixtures["upload_html_new"] = {"status": r.status_code, "body": r.get_json(),
                                   "request": {"html": froala}}

    # upload_html appended into an existing base doc
    ub = upload("uitest.docx")
    r = client.post("/upload_html", json={"html": froala, "append_to": ub["doc_id"]})
    fixtures["upload_html_append"] = {"status": r.status_code, "body": r.get_json(),
                                      "request": {"append_to": ub["doc_id"]}}

    # duplicate a section
    u = upload("uitest.docx")
    r = client.post("/duplicate", json={"doc_id": u["doc_id"], "block_ids": ["1"]})
    fixtures["duplicate"] = {"status": r.status_code, "body": r.get_json(),
                             "request": {"doc_id": u["doc_id"], "block_ids": ["1"]}}

    # find & replace
    u = upload("uitest.docx")
    r = client.post("/find_replace", json={"doc_id": u["doc_id"], "find": "Ann", "replace": "Zoe"})
    fixtures["find_replace"] = {"status": r.status_code, "body": r.get_json(),
                                "request": {"doc_id": u["doc_id"], "find": "Ann", "replace": "Zoe"}}

    # merge popup flow: set_content on the first table with combined+deduped
    # rows, then delete the second table
    u = upload("uitest_merge.docx")
    table_ids = [
        k for k, v in sorted(u["blocks"].items(), key=lambda kv: kv[1]["order"])
        if v["type"] == "table"
    ]
    merged_rows = [["N", "V"], ["a", "1"], ["b", "2"]]  # duplicate header dropped
    r1 = client.post(
        "/table",
        json={"doc_id": u["doc_id"], "block_id": table_ids[0], "operation": "set_content", "rows": merged_rows},
    )
    fixtures["merge_set_content"] = {
        "status": r1.status_code, "body": r1.get_json(),
        "request": {"doc_id": u["doc_id"], "block_id": table_ids[0], "table_ids": table_ids},
    }
    r2 = client.post("/delete", json={"doc_id": u["doc_id"], "block_ids": [table_ids[1]], "cascade": False})
    fixtures["merge_delete"] = {"status": r2.status_code, "body": r2.get_json()}

    return fixtures


if __name__ == "__main__":
    fixtures = capture()
    with open("fixtures.json", "w") as f:
        json.dump(fixtures, f, indent=2)
    print("wrote fixtures.json with keys:", sorted(fixtures))
